"""
Выгрузка справочника в таблицу: блюда по группам, кухни и виды.

Нужна затем, что каталог ведут в админке, а смотреть и обсуждать его удобнее
таблицей. Выгрузка идёт из базы, поэтому разойтись с сайтом ей нечем — в
отличие от таблицы, которую однажды набили руками.
"""

from pathlib import Path

from django.core.management.base import BaseCommand, CommandError

from posts.models import DishGroup, DishType, Taxon

DEFAULT_NAME = 'catalog.xlsx'


def _cuisine_of(dish):
    for taxon in dish.default_taxons.all():
        if taxon.kind == Taxon.KIND_CUISINE:
            return taxon.name
    return '—'


def collect():
    """Три листа: блюда с группой и кухней, кухни, виды."""
    dishes = [('№', 'Группа', 'Блюдо', 'Значок', 'Кухня', 'Иконка')]
    position = 0
    for group in DishGroup.objects.prefetch_related('dish_types__default_taxons'):
        for dish in group.dish_types.all():
            position += 1
            dishes.append((
                position, group.name, dish.name, dish.emoji,
                _cuisine_of(dish), 'да' if dish.icon else 'нет',
            ))
    # Блюдо без группы в интерфейс не попадёт — в таблице оно должно быть
    # видно, иначе пропажу замечают только по жалобе.
    for dish in DishType.objects.filter(group=None).prefetch_related('default_taxons'):
        position += 1
        dishes.append((
            position, 'БЕЗ ГРУППЫ', dish.name, dish.emoji,
            _cuisine_of(dish), 'да' if dish.icon else 'нет',
        ))

    cuisines = [('№', 'Кухня', 'Код', 'Значок', 'Блюд', 'Иконка')]
    for i, taxon in enumerate(Taxon.objects.filter(kind=Taxon.KIND_CUISINE), 1):
        cuisines.append((
            i, taxon.name, taxon.slug, taxon.emoji,
            taxon.dish_types.count(), 'да' if taxon.icon else 'нет',
        ))

    types = [('№', 'Вид', 'Код', 'Значок', 'Иконка')]
    for i, taxon in enumerate(Taxon.objects.filter(kind=Taxon.KIND_TYPE), 1):
        types.append((
            i, taxon.name, taxon.slug, taxon.emoji, 'да' if taxon.icon else 'нет',
        ))

    return [('Блюда', dishes), ('Кухни', cuisines), ('Виды', types)]


class Command(BaseCommand):
    help = 'Выгружает справочник в xlsx (или csv) — блюда, кухни и виды.'

    def add_arguments(self, parser):
        parser.add_argument('--path', default=DEFAULT_NAME, help='Куда сохранить.')
        parser.add_argument(
            '--format', choices=('xlsx', 'csv'), default='xlsx',
            help='csv пишет три файла рядом — по одному на лист.',
        )

    def handle(self, *args, **options):
        sheets = collect()
        path = Path(options['path'])

        if options['format'] == 'csv':
            self._write_csv(path, sheets)
        else:
            self._write_xlsx(path, sheets)

        for name, rows in sheets:
            self.stdout.write(f'  {name}: {len(rows) - 1} строк')
        self.stdout.write(self.style.SUCCESS(f'Готово: {path}'))

    def _write_xlsx(self, path, sheets):
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Alignment, Font, PatternFill
        except ImportError:
            raise CommandError(
                'Нужен openpyxl: pip install openpyxl. '
                'Либо выгрузите в csv: --format csv'
            )

        book = Workbook()
        book.remove(book.active)

        for title, rows in sheets:
            sheet = book.create_sheet(title)
            for row in rows:
                sheet.append(row)

            # Шапку закрепляем и подсвечиваем: без этого на шестидесяти строках
            # приходится всё время вспоминать, что в каком столбце.
            head = PatternFill('solid', fgColor='E9F3EC')
            for cell in sheet[1]:
                cell.font = Font(bold=True)
                cell.fill = head
                cell.alignment = Alignment(vertical='center')
            sheet.freeze_panes = 'A2'
            sheet.auto_filter.ref = sheet.dimensions

            for column in sheet.columns:
                width = max(len(str(c.value or '')) for c in column) + 3
                sheet.column_dimensions[column[0].column_letter].width = min(width, 40)

        book.save(path)

    def _write_csv(self, path, sheets):
        import csv

        for title, rows in sheets:
            target = path.with_name(f'{path.stem}-{title.lower()}.csv')
            # utf-8-sig и точка с запятой: без BOM Excel читает кириллицу
            # кракозябрами, а запятую в русской локали не считает разделителем.
            with open(target, 'w', encoding='utf-8-sig', newline='') as fh:
                csv.writer(fh, delimiter=';').writerows(rows)
            self.stdout.write(f'  записан {target}')
